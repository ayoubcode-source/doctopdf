from flask import Flask, render_template, request, send_file, jsonify
import os
from docx import Document
from docx2pdf import convert

app = Flask(__name__)

# Ensure output directory exists
os.makedirs('output', exist_ok=True)

@app.route('/')
def index():
    return render_template('index.html')

# Function to convert Word document to PDF
def convert_to_pdf(docx_path):
    pdf_path = 'output/output.pdf'
    convert(docx_path, pdf_path)  # Convert using docx2pdf
    return pdf_path

# Generate document and convert to PDF
@app.route('/generate', methods=['POST'])
def generate():
    witness1_name = request.form['witness1_name']
    witness1_id = request.form['witness1_id']
    witness1_address = request.form['witness1_address']
    witness2_name = request.form['witness2_name']
    witness2_id = request.form['witness2_id']
    witness2_address = request.form['witness2_address']
    declaration_content = request.form['declaration_content']
    date = request.form['date']

    # Replace placeholders in the Word template
    modified_doc_path = create_word_from_template(
        witness1_name, witness1_id, witness1_address,
        witness2_name, witness2_id, witness2_address,
        declaration_content, date
    )

    # Convert the modified document to PDF
    pdf_path = convert_to_pdf(modified_doc_path)

    # Return the PDF to download
    return send_file(pdf_path, mimetype='application/pdf')

def create_word_from_template(witness1_name, witness1_id, witness1_address, witness2_name, witness2_id, witness2_address, declaration_content, date):
    # Load the Word template
    doc = Document('اشهاد الشهود.docx')

    # Replace placeholders in paragraphs
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace('[witness1_name]', witness1_name)
            run.text = run.text.replace('[witness1_id]', witness1_id)
            run.text = run.text.replace('[witness1_address]', witness1_address)
            run.text = run.text.replace('[witness_name]', witness2_name)
            run.text = run.text.replace('[witness_id]', witness2_id)
            run.text = run.text.replace('[witness_address]', witness2_address)
            run.text = run.text.replace('[declaration_content]', declaration_content)
            run.text = run.text.replace('[date]', date)

    # Save the modified document
    output_docx_path = 'output/output.docx'
    doc.save(output_docx_path)

    return output_docx_path

if __name__ == '__main__':
    app.run(debug=True)
